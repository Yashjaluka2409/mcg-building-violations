"""Response envelope, pagination, filtering, search and ordering - the DRF conventions the web portal and
the mobile app rely on:

* JSON: datetimes in Asia/Kolkata ISO-8601 (+05:30), decimals as strings, UUIDs as strings
* lists: {count, next, previous, results} with ?page= & ?page_size= (25 default, 500 max)
* ?<field>=value exact filters, ?search= across the declared columns, ?ordering=-field
* CSV / XLSX downloads for the reports (?export=csv|xlsx)
"""
from __future__ import annotations

import csv
import io
import json
import uuid
from datetime import date, datetime
from decimal import Decimal

from fastapi import HTTPException, Request
from fastapi.responses import JSONResponse, Response
from sqlalchemy import inspect, or_
from sqlalchemy.orm import Query
from sqlalchemy.orm.interfaces import MANYTOONE

from app.core.timeutil import aware, localtime


# ---------------------------------------------------------------- JSON
def to_jsonable(v):
    if isinstance(v, dict):
        return {str(k): to_jsonable(x) for k, x in v.items()}
    if isinstance(v, (list, tuple, set)):
        return [to_jsonable(x) for x in v]
    if isinstance(v, datetime):
        return localtime(aware(v)).isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, uuid.UUID):
        return str(v)
    if isinstance(v, bytes):
        return v.decode("utf-8", "replace")
    return v


class BVJSONResponse(JSONResponse):
    def render(self, content) -> bytes:
        return json.dumps(to_jsonable(content), ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode("utf-8")


def resp(data, status: int = 200, headers: dict | None = None) -> BVJSONResponse:
    return BVJSONResponse(content=data, status_code=status, headers=headers)


# ---------------------------------------------------------------- query parameters
def parse_bool(v, default=None):
    if v is None:
        return default
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in ("1", "true", "yes", "on", "t", "y"):
        return True
    if s in ("0", "false", "no", "off", "f", "n", ""):
        return False
    return default


def form_json(v):
    """A JSON object sent as a multipart/form field (the mobile app posts location_integrity this way)."""
    if v is None or v == "":
        return None
    if isinstance(v, (dict, list)):
        return v
    try:
        return json.loads(v)
    except (TypeError, ValueError):
        return None


def _coerce(col, raw: str):
    try:
        pt = col.type.python_type
    except NotImplementedError:
        return raw
    try:
        if pt is bool:
            b = parse_bool(raw)
            if b is None:
                raise ValueError
            return b
        if pt is int:
            return int(raw)
        if pt is uuid.UUID:
            return uuid.UUID(raw)
        if pt is Decimal:
            return Decimal(raw)
        if pt is date:
            return date.fromisoformat(raw)
        if pt is datetime:
            return aware(datetime.fromisoformat(raw))
    except (ValueError, TypeError):
        raise HTTPException(400, f"Invalid value for filter '{col.name}': {raw!r}")
    return raw


def apply_filters(query: Query, model, params, fields) -> Query:
    """DRF/django-filter style exact filters: ?status=DRAFT&zone=2&sla_breached=true."""
    mapper = inspect(model)
    for f in fields:
        raw = params.get(f)
        if raw in (None, ""):
            continue
        if f in mapper.relationships:
            rel = mapper.relationships[f]
            if rel.direction is MANYTOONE:
                col = list(rel.local_columns)[0]
                query = query.filter(col == _coerce(col, raw))
            else:
                target = rel.mapper
                pk = list(target.primary_key)[0]
                query = query.filter(getattr(model, f).any(pk == _coerce(pk, raw)))
        elif f in mapper.columns:
            col = mapper.columns[f]
            query = query.filter(col == _coerce(col, raw))
        elif f + "_id" in mapper.columns:
            col = mapper.columns[f + "_id"]
            query = query.filter(col == _coerce(col, raw))
    return query


def apply_search(query: Query, params, columns) -> Query:
    """?search= : every whitespace-separated term must match (case-insensitively) at least one column."""
    q = (params.get("search") or "").strip()
    if not q or not columns:
        return query
    for term in q.split():
        like = f"%{term}%"
        query = query.filter(or_(*[c.ilike(like) for c in columns]))
    return query


def apply_ordering(query: Query, params, allowed: dict, default) -> Query:
    """?ordering=-created_at,status (only declared fields); `default` is a list of column expressions."""
    raw = (params.get("ordering") or "").strip()
    exprs = []
    for part in [p.strip() for p in raw.split(",") if p.strip()]:
        desc = part.startswith("-")
        name = part.lstrip("-")
        col = allowed.get(name)
        if col is None:
            continue
        exprs.append(col.desc() if desc else col.asc())
    return query.order_by(*(exprs or default))


# ---------------------------------------------------------------- pagination
def paginate(request: Request, query: Query, item_fn, default_size: int = 25, max_size: int = 500) -> dict:
    params = request.query_params
    try:
        page = int(params.get("page", 1))
    except ValueError:
        raise HTTPException(404, "Invalid page.")
    try:
        size = max(1, min(int(params.get("page_size", default_size)), max_size))
    except ValueError:
        size = default_size
    total = query.order_by(None).count()
    if page < 1 or (page > 1 and (page - 1) * size >= total):
        raise HTTPException(404, "Invalid page.")
    rows = query.offset((page - 1) * size).limit(size).all()
    url = request.url
    nxt = str(url.include_query_params(page=page + 1)) if page * size < total else None
    prev = None
    if page > 1:
        prev = str(url.remove_query_params("page")) if page - 1 == 1 else str(url.include_query_params(page=page - 1))
    return {"count": total, "next": nxt, "previous": prev, "results": [item_fn(r) for r in rows]}


# ---------------------------------------------------------------- file downloads
def _cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return localtime(aware(v)).isoformat()
    return v


def csv_response(filename: str, cols: list, rows: list, example: list | None = None) -> Response:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    if example is not None:
        w.writerow(example)
    for r in rows:
        w.writerow([_cell(v) for v in r])
    return Response(content=buf.getvalue(), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def xlsx_response(filename: str, sheet: str, cols: list, rows: list) -> Response:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet[:30]
    ws.append(cols)
    for r in rows:
        ws.append([(localtime(aware(v)).replace(tzinfo=None) if isinstance(v, datetime) else (str(v) if isinstance(v, (uuid.UUID, list, dict)) else v)) for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def read_table_upload(filename: str, raw: bytes) -> list[dict]:
    """CSV or XLSX upload -> list of row dicts keyed by the header row (dates from XLSX as ISO strings)."""
    if (filename or "").lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it, [])]
        rows = []
        for r in it:
            if not any(v not in (None, "") for v in r):
                continue
            rows.append(dict(zip(header, [("" if v is None else (v.date().isoformat() if hasattr(v, "date") and not isinstance(v, str) else v)) for v in r])))
        return rows
    text = raw.decode("utf-8-sig", errors="replace")
    return [dict(r) for r in csv.DictReader(io.StringIO(text))]
